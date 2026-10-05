import io

import pandas as pd
from openpyxl import load_workbook

from fidc import excel


def test_xlsx_com_rotulos_e_formatos():
    df = pd.DataFrame({"nome": ["A"], "pl": [1.5e9], "inad_90": [0.0123], "dt": [pd.Timestamp("2026-08-31")]})
    wb = load_workbook(io.BytesIO(excel.to_xlsx({"Teste": df})))
    ws = wb["Teste"]
    assert [c.value for c in ws[1]] == ["Nome", "PL", "Inad. >90d", "Competência"]
    assert ws["C2"].number_format == "0.00%"
    assert "Fonte" in wb.sheetnames


def test_json_safe_remove_nan():
    df = pd.DataFrame({"a": [1.0, float("nan")], "dt": [pd.Timestamp("2026-01-31"), pd.NaT]})
    assert excel.json_safe(df) == [{"a": 1.0, "dt": "2026-01-31"}, {"a": None, "dt": None}]
