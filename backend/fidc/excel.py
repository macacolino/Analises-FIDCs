"""Export para Excel com rótulos e formatos do dicionário."""
from __future__ import annotations

import io
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import dicionario

NUMFMT = {
    "brl": '#,##0;[Red]-#,##0',
    "pct": "0.00%",
    "pct100": '0.00"%"',
    "x": '0.00"x"',
    "dias": "#,##0",
    "int": "#,##0",
    "data": "yyyy-mm-dd",
}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_sheet(ws, df: pd.DataFrame, nota: str | None = None) -> None:
    row0 = 1
    if nota:
        ws.cell(1, 1, nota).font = Font(italic=True, color="666666")
        row0 = 3
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row0, j, dicionario.label(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), row0 + 1):
        for j, v in enumerate(rec, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
                continue
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            ws.cell(i, j, v)
    for j, col in enumerate(df.columns, 1):
        f = NUMFMT.get(dicionario.fmt(col))
        letter = get_column_letter(j)
        if f:
            for cell in ws[letter][row0:]:
                cell.number_format = f
        width = max(len(dicionario.label(col)), *(len(str(x)) for x in df[col].head(200).tolist() or [""]))
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = ws.cell(row0 + 1, 1)
    ws.auto_filter.ref = f"A{row0}:{get_column_letter(max(1, len(df.columns)))}{row0 + len(df)}"


def to_xlsx(sheets: dict[str, pd.DataFrame], notas: dict[str, str] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        ws = wb.create_sheet(name[:31])
        _write_sheet(ws, df, (notas or {}).get(name))
    ws = wb.create_sheet("Fonte")
    ws["A1"] = "Fonte: CVM - Informe Mensal de FIDC (dados.cvm.gov.br) e cadastro de fundos; FNET (B3)."
    ws["A2"] = f"Gerado em {datetime.now():%Y-%m-%d %H:%M}. Dados declarados pelos administradores, sem auditoria."
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def json_safe(df: pd.DataFrame) -> list[dict]:
    out = df.astype(object).where(pd.notna(df), None).to_dict("records")
    for r in out:
        for k, v in r.items():
            if isinstance(v, (pd.Timestamp, datetime, date)):
                r[k] = v.isoformat()[:10]
            elif isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
                r[k] = None
    return out
